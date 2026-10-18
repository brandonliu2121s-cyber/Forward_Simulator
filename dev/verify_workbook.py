"""Cross-check the workbook against the pure-Python reference model.

    python3 verify_workbook.py                       # check the instrument currently selected in the file
    python3 verify_workbook.py --all <recalc.py>     # switch to EVERY instrument, recalculate with LibreOffice, compare
    python3 verify_workbook.py --all <recalc.py> --snapshot tests/fixtures/snapshot_fixture.json
                                                     # same, but build with a live snapshot and Use_Live_Spot = 1

<recalc.py> is the xlsx skill's recalculation script (LibreOffice headless).
"""
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "python"))
from instruments import ALL, params
from live_data import read_snapshot
from reference_model import fair_forward, simulate

PATH = os.path.join(HERE, "..", "excel", "Forwards_Market_Making_Simulator.xlsx")
COLS = dict(spot="D", fwd="E", pos_open="G", skew="I", bid="J", ask="K", fill="M", edge="P", mtm="Q",
            pos_breach="S", loss_breach="T", hedge="V", hedge_cost="X", pos_close="Y",
            step_pnl="Z", cum="AA", halted="AB")


def check(path, instrument, snapshot=None):
    wb = load_workbook(path, data_only=True)
    ws = wb["Simulation"]
    assert wb["Inputs"]["C4"].value == instrument, "selector mismatch"
    p = params(instrument)
    lib_spot = p["spot_init"]
    live = ((snapshot or {}).get("quotes", {}).get(instrument) or {}).get("price") if snapshot else None
    if snapshot:                       # workbook was built with Use_Live_Spot = 1
        flag = wb["Inputs"][wb.defined_names["Use_Live_Spot"].attr_text.split("!")[1].replace("$", "")].value
        assert flag == 1, "Use_Live_Spot should be 1"
        if live:
            p["spot_init"] = live      # a missing / failed quote must fall back to the library spot
    zs = [ws.cell(row=r, column=2).value for r in range(7, 67)]
    acts = [ws.cell(row=r, column=3).value for r in range(7, 67)]
    ref = simulate(zs, acts, p)
    bad = 0
    for i, row in enumerate(ref):
        for key, col in COLS.items():
            x = ws[f"{col}{7 + i}"].value
            tol = 1e-6 * max(1.0, abs(row[key]))
            if abs((x or 0) - row[key]) > tol:
                bad += 1
                if bad <= 5:
                    print(f"  MISMATCH {instrument} step {i + 1} {key}: excel={x} ref={row[key]}")
    # fair forward in the Instruments library table
    names = [c.value for c in wb["Instruments"]["B"][5:5 + len(ALL)]]
    row = 6 + names.index(instrument)
    fx = wb["Instruments"][f"V{row}"].value
    ref_fx = fair_forward(lib_spot, p)
    if abs(fx - ref_fx) > 1e-9 * lib_spot:
        bad += 1
        print(f"  MISMATCH {instrument} Instruments fair forward: excel={fx}")
    if snapshot:                       # Live sheet: fair forward at the live spot (n/a when no quote)
        lv = wb["Live"][f"I{row}"].value
        want = fair_forward(live, p) if live else "n/a"
        if (want == "n/a" and lv != "n/a") or (want != "n/a" and abs(lv - want) > 1e-9 * live):
            bad += 1
            print(f"  MISMATCH {instrument} Live fair forward: excel={lv} ref={want}")
    # Summary integrity checks
    su = wb["Summary"]
    checks = [su.cell(row=r, column=3).value for r in range(1, 60) if su.cell(row=r, column=3).value in ("OK", "ERROR")]
    if checks != ["OK"] * 3:
        bad += 1
        print(f"  INTEGRITY CHECKS {instrument}: {checks}")
    print(f"{instrument:18s} {len(ref) * len(COLS)} cells checked, {bad} problems")
    return bad


if __name__ == "__main__":
    if "--all" in sys.argv:
        recalc = sys.argv[sys.argv.index("--all") + 1]
        snap_path = sys.argv[sys.argv.index("--snapshot") + 1] if "--snapshot" in sys.argv else None
        snapshot = read_snapshot(snap_path) if snap_path else None
        total = 0
        with tempfile.TemporaryDirectory() as tmp:
            base = PATH
            if snapshot:               # build a workbook that contains the snapshot
                base = str(Path(tmp) / "base_live.xlsx")
                subprocess.run([sys.executable, os.path.join(HERE, "build_simulator.py"), "--snapshot", snap_path, "--out", base],
                               check=True, capture_output=True)
            for inst in ALL:
                f = Path(tmp) / "wb.xlsx"
                wb = load_workbook(base)
                wb["Inputs"]["C4"] = inst["name"]
                if snapshot:
                    wb["Inputs"][wb.defined_names["Use_Live_Spot"].attr_text.split("!")[1].replace("$", "")] = 1
                wb.save(f)
                out = subprocess.run([sys.executable, recalc, str(f), "120"], capture_output=True, text=True).stdout
                res = json.loads(out)
                if res.get("status") != "success":
                    print(inst["name"], "RECALC ERRORS:", res)
                    total += 1
                    continue
                total += check(f, inst["name"], snapshot)
        print("TOTAL PROBLEMS:", total)
        sys.exit(1 if total else 0)
    wb = load_workbook(PATH, data_only=True)
    sys.exit(1 if check(PATH, wb["Inputs"]["C4"].value) else 0)
