"""Builds Forwards_Market_Making_Simulator.xlsx: a multi-instrument forwards market-making simulation.

Pick any FX pair or commodity on the Inputs sheet; the whole workbook re-prices. Every number it
computes is a live Excel formula driven by the Instruments table, the Inputs sheet and the two blue
per-step input columns on the Simulation sheet (market move, client action).

Run:  python3 build_simulator.py   then   python3 <skill>/scripts/recalc.py Forwards_Market_Making_Simulator.xlsx
"""
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

import argparse
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "python"))
from instruments import ALL, DEFAULT, GLOBALS
from live_data import read_snapshot
from reference_model import generate_scenario

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "excel", "Forwards_Market_Making_Simulator.xlsx")
SEED, N_STEPS = 277, 60

_ap = argparse.ArgumentParser()
_ap.add_argument("--snapshot", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "live_snapshot.json"),
                 help="live spot snapshot written by python/refresh_live.py (default: live_snapshot.json in the repo root)")
_ap.add_argument("--out", default=OUT, help="output .xlsx path")
_ARGS = _ap.parse_args()
# None if absent: the Live sheet is then empty and the toggle falls back to library spots
SNAP = read_snapshot(_ARGS.snapshot)

FONT = "Arial"
NAVY, GREY, LIGHT = "1F3864", "F2F2F2", "D9E1F2"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_sub = Font(name=FONT, size=10, italic=True, color="595959")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_group = PatternFill("solid", fgColor=LIGHT)
fill_input = PatternFill("solid", fgColor="FFFF99")
fill_grey = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

NF_PX = "#,##0.0000"
NF_QTY = "#,##0;(#,##0);-"
NF_PNL = "#,##0.00;(#,##0.00);-"
NF_PCT = "0.0%;(0.0%);-"
NF_TICK = "0.00;(0.00);-"
NF_PIP = "#,##0.0;(#,##0.0);-"

wb = Workbook()


def put(ws, ref, value, font=f_base, nf=None, fill=None, align=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if nf:
        c.number_format = nf
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def band(ws, row, label, cols, font=f_head, fill=fill_head):
    """Section band across the given column letters (label in the first)."""
    for i, col in enumerate(cols):
        put(ws, f"{col}{row}", label if i == 0 else None, font, fill=fill)


# --------------------------------------------------------------------------------------- Guide
g = wb.active
g.title = "Guide"
g.sheet_view.showGridLines = False
widths(g, {"A": 2, "B": 26, "C": 105})
put(g, "B1", "Forwards Market-Making Simulator", f_title)
put(g, "B2", "Pick any FX pair or commodity: the model prices its forward, quotes a two-way market with inventory skew, "
             "and hedges back-to-back when position or loss limits are breached.", f_sub)
rows = [
    ("START HERE", None),
    ("1. Choose an instrument", "Inputs!C4 is a dropdown of 7 FX pairs, 6 commodities and one Custom row. Everything (pricing, "
                                "quote size, spread, limits, labels, charts) re-prices when you change it."),
    ("2. Tweak assumptions", "Inputs shows each instrument's default, an optional Override column (blue on yellow) and the ACTIVE "
                             "value the model uses. Leave Override blank to use the default."),
    ("3. Run a scenario", "On Simulation, edit the blue columns: market move (in step-vol units, so one scenario works for "
                          "every instrument) and client action (BUY / SELL / NONE)."),
    ("4. Read the results", "Summary: fills, hedges, P&L attribution, limit usage, checks and charts. Pricing: quote calculator "
                            "and skew ladder for any inventory."),
    ("SHEETS", None),
    ("Inputs", "Instrument selector, defaults + overrides, global risk settings."),
    ("Instruments", "The instrument library (one row each) with a live fair forward, forward points and premium/discount for "
                    "every instrument. Edit a row or overwrite the Custom row to add your own."),
    ("Pricing", "Step-by-step derivation of the fair forward for the selected instrument, a live quote calculator and a skew ladder."),
    ("Simulation", "60-step market replay: market move, client RFQ, quotes, fill, P&L, limit checks and hedge for every step."),
    ("Live", "Optional live-market snapshot (spots only) written by refresh_live.py, with the fair forward at the live spot for every instrument."),
    ("Summary", "Fills, hedges, P&L attribution (spread capture / mark-to-market / hedge cost), limit usage, checks, charts."),
    ("LIVE DATA (OPTIONAL)", None),
    ("How to use it", "Run  python3 refresh_live.py  on a computer with internet (see README). It fetches spot prices into live_snapshot.json and "
                      "rebuilds this workbook. Then set Inputs > 'Use live spot' = 1. Default is 0, so the workbook never depends on a feed."),
    ("What is and is not live", "Only SPOT prices come from the feed (a snapshot, not a stream; free sources are delayed). Interest rates, convenience yields, "
                                "sizes and limits stay manual on the Instruments sheet. Commodity prices are front-month futures, not spot."),
    ("PRICING", None),
    ("One formula, two readings", "F = Spot x (1 + rate B x d / basis B) / (1 + rate A x d / basis A)."),
    ("FX pair BASE/QUOTE", "Covered interest parity: rate A = BASE-currency rate, rate B = QUOTE-currency rate. Higher-yielding "
                           "currency trades at a forward discount. Position unit = BASE currency; P&L in QUOTE currency."),
    ("Commodity", "Cost of carry: rate A = convenience yield, rate B = financing + storage cost. Carry > yield gives contango "
                  "(forward above spot, e.g. Gold, Natural Gas); yield > carry gives backwardation (e.g. WTI). P&L in USD."),
    ("QUOTING AND RISK", None),
    ("Quotes", "Mid = fair forward + skew. Bid = round(fair/tick + skew - spread/2) x tick, Ask = Bid + spread. Fixed clip size each side."),
    ("Skew", "Skew (ticks) = -(position / position limit, capped at +/-100%) x max skew. Long inventory pushes both quotes down "
             "(cheap to buy from us, unattractive to sell to us); short pushes them up."),
    ("Position limit", "If |position| after a client fill exceeds the limit, a back-to-back hedge trades the street back to the "
                       "residual (% of limit, default 0 = flat), paying the street half-spread."),
    ("Loss limit", "If cumulative P&L (before that step's hedge cost) reaches -loss limit, the book is flattened. With 'halt' = 1 "
                   "quoting then stops for the rest of the run."),
    ("CONVENTIONS", None),
    ("Position", "Net units bought forward from clients (+ = long). Client BUY lifts our Ask (we go shorter); client SELL hits our Bid."),
    ("P&L", "In the quote/P&L currency, measured against fair forward: spread capture + mark-to-market of opening position on the "
            "fair-forward move + street hedge cost."),
    ("Market move", "Spot change per step = move (sigma units) x step vol (ticks) x tick. Step vol is set per instrument."),
    ("LIMITATIONS", None),
    ("Simplifications", "Constant tenor each step (no roll-down); P&L not discounted; one RFQ per step at a fixed size; hedges fill "
                        "instantly at fair +/- street half-spread. ALL spots, rates, sizes and limits are illustrative placeholders, "
                        "not live market data: update them before drawing conclusions. The default scenario is randomly generated "
                        "(seed %d) for demonstration." % SEED),
]
r = 4
for k, v in rows:
    if v is None:
        band(g, r, k, "BC")
    else:
        put(g, f"B{r}", k, f_bold, align=Alignment(vertical="top", wrap_text=True))
        put(g, f"C{r}", v, align=Alignment(wrap_text=True, vertical="top"))
        g.row_dimensions[r].height = 14 * (1 + len(v) // 108)
    r += 1

# --------------------------------------------------------------------------------------- Instruments
it = wb.create_sheet("Instruments")  # created before Inputs refs; reordered below
it.sheet_view.showGridLines = False
IF_, IL_ = 6, 6 + len(ALL) - 1
put(it, "B1", "Instrument Library", f_title)
put(it, "B2", "One row per instrument. Blue = editable input; black = formula. Spots, rates, sizes and limits are ILLUSTRATIVE "
              "placeholders, not live market data. Overwrite the Custom row (yellow) to add your own.", f_sub)
put(it, "B3", "Rate A / Rate B: FX = base / quote currency interest rate; Commodity = convenience yield / (financing + storage cost). "
              "Loss limit default = 6 x position limit x step vol x tick.", f_sub)
icols = [  # key, header, col letter, width, number format
    ("name", "Instrument", "B", 17, None), ("type", "Type", "C", 11, None), ("unit", "Position unit", "D", 9, None),
    ("pnl", "P&L ccy", "E", 8, None), ("spot", "Spot", "F", 12, NF_PX), ("rate_a", "Rate A", "G", 9, "0.00%"),
    ("financing", "Financing (commodity)", "H", 12, "0.00%"), ("storage", "Storage (commodity)", "I", 12, "0.00%"),
    ("rate_b", "Rate B", "J", 9, "0.00%"), ("basis_a", "Day basis A", "K", 8, "0"), ("basis_b", "Day basis B", "L", 8, "0"),
    ("tenor_days", "Tenor (days)", "M", 8, "0"), ("tick", "Tick size", "N", 9, "0.0000"),
    ("clip", "Quote size (units)", "O", 11, NF_QTY), ("spread_ticks", "Spread (ticks)", "P", 8, "0.0"),
    ("max_skew_ticks", "Max skew (ticks)", "Q", 8, "0.0"), ("pos_limit", "Position limit (units)", "R", 12, NF_QTY),
    ("loss_limit", "Loss limit (P&L ccy)", "S", 12, NF_QTY), ("hedge_cost_ticks", "Hedge cost (ticks)", "T", 8, "0.0"),
    ("step_vol_ticks", "Step vol (ticks)", "U", 8, "0.0"),
]
CALC = [("V", "Fair forward", 13, NF_PX), ("W", "Fwd points (ticks)", 11, NF_PIP), ("X", "Premium / (discount) p.a.", 12, "0.00%;(0.00%)")]
widths(it, {"A": 2, "Y": 110})
for key, head, col, w, nf in icols:
    it.column_dimensions[col].width = w
    put(it, f"{col}5", head, f_head, fill=fill_head, align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=box)
for col, head, w, nf in CALC:
    it.column_dimensions[col].width = w
    put(it, f"{col}5", head, f_head, fill=fill_head, align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=box)
put(it, "Y5", "Source / note", f_head, fill=fill_head, align=Alignment(vertical="center"))
it.row_dimensions[5].height = 42
for k, inst in enumerate(ALL):
    r = IF_ + k
    custom = inst["name"].startswith("Custom")
    for key, head, col, w, nf in icols:
        v = inst[key]
        if key == "rate_b" and inst["type"] == "Commodity":
            put(it, f"{col}{r}", f"=H{r}+I{r}", f_base, nf, border=box)       # financing + storage
        elif v is None:
            put(it, f"{col}{r}", "n/a", f_sub, align=Alignment(horizontal="center"), border=box)
        else:
            put(it, f"{col}{r}", v, f_input, nf, fill_input if custom else None, border=box,
                align=Alignment(horizontal="center") if key in ("type", "unit", "pnl") else None)
    put(it, f"V{r}", f"=F{r}*(1+J{r}*M{r}/L{r})/(1+G{r}*M{r}/K{r})", f_base, NF_PX, border=box)
    put(it, f"W{r}", f"=(V{r}-F{r})/N{r}", f_base, NF_PIP, border=box)
    put(it, f"X{r}", f"=(V{r}/F{r}-1)*365/M{r}", f_base, "0.00%;(0.00%)", border=box)
    put(it, f"Y{r}", inst["note"], f_sub)
it.conditional_formatting.add(f"B{IF_}:X{IL_}", FormulaRule(formula=[f"$B{IF_}=Instrument"], fill=PatternFill("solid", bgColor="BDD7EE")))
it.freeze_panes = "C6"
put(it, f"B{IL_ + 2}", "Highlighted row = instrument currently selected on Inputs.", f_sub)

# --------------------------------------------------------------------------------------- Live
lv = wb.create_sheet("Live")
lv.sheet_view.showGridLines = False
widths(lv, {"A": 2, "B": 18, "C": 14, "D": 30, "E": 26, "F": 28, "G": 14, "H": 14, "I": 15, "J": 14})
put(lv, "B1", "Live Market Snapshot", f_title)
put(lv, "B2", ("Snapshot fetched (UTC): %s" % SNAP["fetched_utc"]) if SNAP else
    "No snapshot loaded: run  python3 refresh_live.py  to populate this sheet, then rebuild.", f_sub)
put(lv, "B3", "Written by refresh_live.py (blue = pasted data). Delayed / unofficial data, educational use only. Only SPOT is live: "
              "rates stay manual on the Instruments sheet.", f_sub)
for col, h in zip("BCDEFGHIJ", ["Instrument", "Live spot", "As of", "Source", "Status", "Library spot", "Live vs library",
                                 "Fair forward at live spot", "Fwd points at live spot (ticks)"]):
    put(lv, f"{col}5", h, f_head, fill=fill_head, border=box,
        align=Alignment(horizontal="center", vertical="center", wrap_text=True))
lv.row_dimensions[5].height = 42
for k, inst in enumerate(ALL):
    r = IF_ + k
    q = (SNAP or {}).get("quotes", {}).get(inst["name"], {})
    put(lv, f"B{r}", f"=Instruments!B{r}", f_link, border=box)
    put(lv, f"C{r}", q.get("price"), f_input, NF_PX, border=box)
    put(lv, f"D{r}", q.get("as_of") or "", f_input, border=box)
    put(lv, f"E{r}", q.get("source") or "", f_input, border=box)
    put(lv, f"F{r}", q.get("status") or "no snapshot", f_input, border=box)
    put(lv, f"G{r}", f"=Instruments!F{r}", f_link, NF_PX, border=box)
    put(lv, f"H{r}", f'=IF(ISNUMBER(C{r}),C{r}/G{r}-1,"n/a")', f_base, "0.0%;(0.0%)", border=box, align=Alignment(horizontal="right"))
    put(lv, f"I{r}", f'=IF(ISNUMBER(C{r}),C{r}*(1+Instruments!J{r}*Instruments!M{r}/Instruments!L{r})/(1+Instruments!G{r}*Instruments!M{r}/Instruments!K{r}),"n/a")',
        f_base, NF_PX, border=box, align=Alignment(horizontal="right"))
    put(lv, f"J{r}", f'=IF(ISNUMBER(C{r}),(I{r}-C{r})/Instruments!N{r},"n/a")', f_base, NF_PIP, border=box, align=Alignment(horizontal="right"))
lv.conditional_formatting.add(f"B{IF_}:J{IL_}", FormulaRule(formula=[f"$B{IF_}=Instrument"], fill=PatternFill("solid", bgColor="BDD7EE")))
lv.freeze_panes = "C6"

# --------------------------------------------------------------------------------------- Inputs
ws = wb.create_sheet("Inputs")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 2, "B": 52, "C": 20, "D": 20, "E": 18, "F": 16, "G": 88})
put(ws, "B1", "Inputs & Assumptions", f_title)
put(ws, "B2", "Blue on yellow = editable. Pick an instrument; leave Override blank to use its default. "
              "The ACTIVE column (bold) is what the model uses.", f_sub)
put(ws, "B4", "Select instrument  (FX pair or commodity)", f_bold)
put(ws, "C4", DEFAULT, Font(name=FONT, size=11, bold=True, color="0000FF"), fill=fill_input, border=box,
    align=Alignment(horizontal="center"))
put(ws, "D4", "<- dropdown; change me", f_sub)
name("Instrument", "Inputs!$C$4")
put(ws, "B5", "Row in Instruments table")
put(ws, "C5", f"=MATCH(Instrument,Instruments!$B${IF_}:$B${IL_},0)", f_base, "0", align=Alignment(horizontal="center"))
name("Inst_Row", "Inputs!$C$5")
dv_inst = DataValidation(type="list", formula1=f"=Instruments!$B${IF_}:$B${IL_}", allow_blank=False,
                         showErrorMessage=True, errorTitle="Instrument", error="Pick an instrument from the list.")
ws.add_data_validation(dv_inst)
dv_inst.add("C4")
put(ws, "B6", "Instrument note")
put(ws, "C6", f"=INDEX(Instruments!$Y${IF_}:$Y${IL_},Inst_Row)", f_sub)
put(ws, "B7", "Spot source in use")
LIVE_OK = f"AND(Use_Live_Spot=1,ISNUMBER(INDEX(Live!$C${IF_}:$C${IL_},Inst_Row)))"
put(ws, "C7", f'=IF({LIVE_OK},"LIVE spot, as of "&INDEX(Live!$D${IF_}:$D${IL_},Inst_Row)&"  ["&INDEX(Live!$E${IF_}:$E${IL_},Inst_Row)&"]",'
    '"Library spot (illustrative placeholder)")', f_bold)

for col, h in zip("BCDEFG", ["Assumption", "Instrument default", "Override (optional)", "ACTIVE value", "Unit", "Note"]):
    put(ws, f"{col}8", h, f_head, fill=fill_head, align=Alignment(horizontal="center" if col in "CDEF" else None))

USD_PER = '=PnL_Ccy&" per "&Unit_Label'
spec = [
    ("section", "INSTRUMENT"),
    ("text", "Inst_Type", "Instrument type", "C", "FX = interest-rate parity; Commodity = cost of carry."),
    ("text", "Unit_Label", "Position unit", "D", "FX: base currency. Commodity: contract unit (oz, bbl, t, MMBtu)."),
    ("text", "PnL_Ccy", "P&L / quote currency", "E", "All P&L, loss limit and hedge cost are in this currency."),
    ("section", "PRICING"),
    ("num", "Spot_Init", "Starting spot price", "F", NF_PX, USD_PER,
     "Default = live spot (Live sheet) when 'Use live spot' = 1 and a snapshot quote exists, else the library spot (illustrative)."),
    ("global", "Use_Live_Spot", "Use live spot from the Live sheet  (1 = yes, 0 = no)", 0, "0", "flag",
     "Run refresh_live.py first. 0 = always use the library spot, so the workbook never depends on a feed."),
    ("num", "Rate_A", '=IF(Inst_Type="FX","Rate A: base-currency interest rate","Rate A: convenience yield")', "G", "0.00%", "% p.a.",
     "Denominator of the forward formula. USD/SGD brief: USD 4.00%."),
    ("num", "Rate_B", '=IF(Inst_Type="FX","Rate B: quote-currency interest rate","Rate B: financing + storage cost")', "J", "0.00%", "% p.a.",
     "Numerator of the forward formula. USD/SGD brief: SGD 0.55%. Commodity: financing + storage."),
    ("num", "Basis_A", "Day-count basis for rate A", "K", "0", "days", "Money-market convention: ACT/360 for USD, EUR, JPY, CHF; ACT/365 for GBP, AUD, SGD, CNH."),
    ("num", "Basis_B", "Day-count basis for rate B", "L", "0", "days", ""),
    ("num", "Tenor_Days", "Forward tenor", "M", "0", "days", "USD/SGD brief: 45 days."),
    ("section", "QUOTING"),
    ("num", "Tick", "Tick size", "N", "0.0000", USD_PER, "1 tick = smallest quoted price increment."),
    ("num", "Clip_Size", "Quote size per side", "O", NF_QTY, "=Unit_Label", "USD/SGD brief: USD 250k per side."),
    ("num", "Spread_Ticks", "Bid/offer spread", "P", "0.0", "ticks", "USD/SGD brief: 5 ticks."),
    ("num", "Max_Skew_Ticks", "Max skew at full position limit", "Q", "0.0", "ticks", "How hard quotes lean away from inventory at limit; linear below."),
    ("section", "RISK LIMITS & HEDGING"),
    ("num", "Pos_Limit", "Position limit (absolute)", "R", NF_QTY, "=Unit_Label", "Breach -> back-to-back hedge with the street."),
    ("num", "Loss_Limit", "Cumulative loss limit", "S", NF_QTY, "=PnL_Ccy", "Breach -> flatten the book with the street."),
    ("num", "Hedge_Cost_Ticks", "Street half-spread paid on hedges", "T", "0.0", "ticks", "Hedges execute at fair forward -/+ this many ticks."),
    ("global", "Hedge_Target_Pct", "Residual position after limit hedge (% of position limit)", GLOBALS["hedge_target_pct"], "0%", "% of limit",
     "0% = hedge fully flat; 25% keeps a quarter-limit residual in the same direction. Assumption."),
    ("derived", "Hedge_Target", "  = residual position", "=Hedge_Target_Pct*Pos_Limit", NF_QTY, "=Unit_Label", ""),
    ("global", "Halt_On_Loss", "Halt quoting after loss-limit breach  (1 = yes, 0 = no)", GLOBALS["halt_on_loss"], "0", "flag",
     "0 = keep quoting (the desk re-flattens every step cumulative P&L stays below the limit). Assumption."),
    ("section", "SIMULATION"),
    ("num", "Step_Vol_Ticks", "Step volatility (1.0 sigma move per step)", "U", "0.0", "ticks", "Simulation market moves are in sigma units x this x tick."),
    ("global", "Start_Pos", "Opening position", GLOBALS["start_pos"], NF_QTY, "=Unit_Label", "Assumption: start flat."),
]
r = 9
for item in spec:
    kind = item[0]
    if kind == "section":
        band(ws, r, item[1], "BCDEFG", f_bold, fill_group)
    elif kind in ("text", "num"):
        nm, label, col = item[1], item[2], item[3]
        put(ws, f"B{r}", label)
        put(ws, f"C{r}", f"=INDEX(Instruments!${col}${IF_}:${col}${IL_},Inst_Row)", f_base,
            item[4] if kind == "num" else None, align=Alignment(horizontal="right"))
        if nm == "Spot_Init":      # live-aware default
            ws[f"C{r}"].value = (f"=IF({LIVE_OK},INDEX(Live!$C${IF_}:$C${IL_},Inst_Row),"
                                 f"INDEX(Instruments!${col}${IF_}:${col}${IL_},Inst_Row))")
        if kind == "num":
            put(ws, f"D{r}", None, f_input, item[4], fill_input, Alignment(horizontal="right"), box)
            put(ws, f"E{r}", f'=IF(D{r}="",C{r},D{r})', f_bold, item[4], align=Alignment(horizontal="right"))
            put(ws, f"F{r}", item[5], f_sub, align=Alignment(horizontal="center"))
            put(ws, f"G{r}", item[6], f_sub)
        else:
            put(ws, f"D{r}", "-", f_sub, align=Alignment(horizontal="center"))
            put(ws, f"E{r}", f"=C{r}", f_bold, align=Alignment(horizontal="right"))
            put(ws, f"G{r}", item[4], f_sub)
        name(nm, f"Inputs!$E${r}")
    elif kind == "global":
        nm, label, val, nf, unit, note = item[1:]
        put(ws, f"B{r}", label)
        for col in "CD":
            put(ws, f"{col}{r}", "-", f_sub, align=Alignment(horizontal="center"))
        put(ws, f"E{r}", val, f_input, nf, fill_input, Alignment(horizontal="right"), box)
        put(ws, f"F{r}", unit, f_sub, align=Alignment(horizontal="center"))
        put(ws, f"G{r}", note, f_sub)
        name(nm, f"Inputs!$E${r}")
        if nm in ("Halt_On_Loss", "Use_Live_Spot"):
            dv_flag = DataValidation(type="list", formula1='"0,1"', allow_blank=False)
            ws.add_data_validation(dv_flag)
            dv_flag.add(f"E{r}")
    elif kind == "derived":
        nm, label, fml, nf, unit, note = item[1:]
        put(ws, f"B{r}", label)
        for col in "CD":
            put(ws, f"{col}{r}", "-", f_sub, align=Alignment(horizontal="center"))
        put(ws, f"E{r}", fml, f_bold, nf, align=Alignment(horizontal="right"))
        put(ws, f"F{r}", unit, f_sub, align=Alignment(horizontal="center"))
        name(nm, f"Inputs!$E${r}")
    r += 1
ws.freeze_panes = "A9"

# --------------------------------------------------------------------------------------- Pricing
pr = wb.create_sheet("Pricing")
pr.sheet_view.showGridLines = False
widths(pr, {"A": 2, "B": 50, "C": 18, "D": 14, "E": 14, "F": 14, "G": 24, "H": 4})
put(pr, "B1", '="Pricing: "&Instrument&" ("&Inst_Type&")"', f_title)
put(pr, "B2", "Fair forward from interest-rate parity (FX) or cost of carry (commodity), then the skewed two-way quote for any inventory.", f_sub)

band(pr, 4, "1. Fair forward", "BCD")
calc = [
    ('="Spot ("&PnL_Ccy&" per "&Unit_Label&")"', "=Spot_Init", NF_PX, f_link),
    ('=IF(Inst_Type="FX","Rate-A factor  1 + r_base x d / basis","Convenience-yield factor  1 + y x d / basis")',
     "=1+Rate_A*Tenor_Days/Basis_A", "0.000000", f_base),
    ('=IF(Inst_Type="FX","Rate-B factor  1 + r_quote x d / basis","Carry factor  1 + (financing + storage) x d / basis")',
     "=1+Rate_B*Tenor_Days/Basis_B", "0.000000", f_base),
    ("Fair forward outright  =  Spot x B-factor / A-factor", "=C5*C7/C6", NF_PX, f_bold),
    ("Forward points (ticks)", "=(C8-C5)/Tick", NF_PIP, f_base),
    ("Forward premium / (discount) p.a.", "=(C8/C5-1)*365/Tenor_Days", "0.00%;(0.00%)", f_base),
    ("Rate differential  (B - A)", "=Rate_B-Rate_A", "0.00%;(0.00%)", f_base),
]
for i, (lab, fml, nf, fnt) in enumerate(calc):
    put(pr, f"B{5 + i}", lab)
    put(pr, f"C{5 + i}", fml, fnt, nf, align=Alignment(horizontal="right"))
put(pr, "D8", '=IF(C8<C5,"Forward at a DISCOUNT to spot"&IF(Inst_Type="FX"," (base-ccy yield > quote-ccy yield)"," (backwardation: yield > carry)"),'
              '"Forward at a PREMIUM to spot"&IF(Inst_Type="FX"," (quote-ccy yield > base-ccy yield)"," (contango: carry > yield)"))', f_sub)

band(pr, 13, "2. Live quote calculator", "BCD")
put(pr, "B14", "Current position, in clips  (edit me; + = long)")
put(pr, "C14", 2, f_input, "0.0;(0.0);0", fill_input, Alignment(horizontal="right"), box)
q = [
    ('="Current position ("&Unit_Label&")"', "=C14*Clip_Size", NF_QTY),
    ("Position utilisation of limit", "=C15/Pos_Limit", NF_PCT),
    ("Skew (ticks)", "=-MAX(-1,MIN(1,C16))*Max_Skew_Ticks", NF_TICK),
    ("Bid  (client sells to us)", "=ROUND(C8/Tick+C17-Spread_Ticks/2,0)*Tick", NF_PX),
    ("Ask  (client buys from us)", "=C18+Spread_Ticks*Tick", NF_PX),
    ("Quoted spread (ticks)", "=(C19-C18)/Tick", NF_TICK),
    ("Mid vs fair forward (ticks)", "=((C18+C19)/2-C8)/Tick", NF_TICK),
    ('="Quote size each side ("&Unit_Label&")"', "=Clip_Size", NF_QTY),
    ("Position after being lifted on the Ask", "=C15-Clip_Size", NF_QTY),
    ("Position after being hit on the Bid", "=C15+Clip_Size", NF_QTY),
    ("Next fill triggers hedge?", '=IF(OR(ABS(C23)>Pos_Limit,ABS(C24)>Pos_Limit),"YES: a fill can breach the limit","No")', None),
]
for i, (lab, fml, nf) in enumerate(q):
    put(pr, f"B{15 + i}", lab)
    put(pr, f"C{15 + i}", fml, f_bold if i in (3, 4) else f_base, nf, align=Alignment(horizontal="right"))

band(pr, 28, "3. Skew ladder: quotes at every inventory level", "BCDEFG")
for col, h in zip("BCDEFG", ["Position (units)", "Utilisation", "Skew (ticks)", "Bid", "Ask", "Limit check"]):
    put(pr, f"{col}29", h, f_bold, fill=fill_group, align=Alignment(horizontal="center", wrap_text=True))
for i in range(13):
    rr = 30 + i
    put(pr, f"B{rr}", "=-6*Clip_Size" if i == 0 else f"=B{rr - 1}+Clip_Size", f_base, NF_QTY)
    put(pr, f"C{rr}", f"=B{rr}/Pos_Limit", f_base, NF_PCT)
    put(pr, f"D{rr}", f"=-MAX(-1,MIN(1,C{rr}))*Max_Skew_Ticks", f_base, NF_TICK)
    put(pr, f"E{rr}", f"=ROUND($C$8/Tick+D{rr}-Spread_Ticks/2,0)*Tick", f_base, NF_PX)
    put(pr, f"F{rr}", f"=E{rr}+Spread_Ticks*Tick", f_base, NF_PX)
    put(pr, f"G{rr}", f'=IF(ABS(B{rr})>Pos_Limit,"BREACH: hedge","within limit")', f_base, align=Alignment(indent=1))
pr.conditional_formatting.add("B30:G42", FormulaRule(formula=["ABS($B30)>Pos_Limit"], fill=PatternFill("solid", bgColor="F8CBAD")))

# --------------------------------------------------------------------------------------- Simulation
sm = wb.create_sheet("Simulation")
sm.sheet_view.showGridLines = False
FIRST, LAST = 7, 6 + N_STEPS  # step 0 lives on row 6
put(sm, "A1", '="Market-Making Simulation: "&Instrument&" ("&Inst_Type&")"', f_title)
put(sm, "A2", "One client RFQ per step. Edit the blue columns B:C to run your own scenario; all other cells are formulas. "
              "Position = net units bought from clients (+ long).", f_sub)
put(sm, "A3", "Default scenario: market moves and client flow randomly generated by build_simulator.py (seed %d); "
              "illustrative only. Moves are in sigma units, so the same scenario runs on any instrument." % SEED, f_sub)

groups = [("A", "C", "Scenario input"), ("D", "F", "Market"), ("G", "I", "Exposure & skew"), ("J", "L", "Our quote"),
          ("M", "P", "Client trade"), ("Q", "T", "P&L and limit checks"), ("U", "X", "Back-to-back hedge"),
          ("Y", "AC", "Close of step")]
for a, b, label in groups:
    sm.merge_cells(f"{a}4:{b}4")
    put(sm, f"{a}4", label, f_bold, fill=fill_group, align=Alignment(horizontal="center"))
U, P = '"&Unit_Label&"', '"&PnL_Ccy&"'
heads = [
    "Step", "Market move\n(sigma units)", "Client action\n(BUY / SELL / NONE)",
    "Spot", "Fair forward", "Fwd points (ticks)",
    f'="Opening position ({U})"', "Utilisation of limit", "Skew (ticks)",
    "Bid", "Ask", "Quoting status",
    f'="Client fill ({U}, + = we buy)"', "Fill price", "Position after client trade", f'="Spread capture ({P})"',
    f'="MTM on opening position ({P})"', f'="Cum. P&L before hedge ({P})"', "Position limit breach", "Loss limit breach",
    "Hedge trigger", f'="Hedge qty ({U}, + = we buy)"', "Hedge price", f'="Hedge cost ({P})"',
    f'="Closing position ({U})"', f'="Step P&L ({P})"', f'="Cumulative P&L ({P})"', "Quoting halted (1 = yes)",
    f'="Closing position value ({P})"',
]
for i, h in enumerate(heads, start=1):
    put(sm, f"{get_column_letter(i)}5", h, f_head, fill=fill_head,
        align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=box)
sm.row_dimensions[5].height = 54
widths(sm, {get_column_letter(i): w for i, w in enumerate(
    [6, 11, 15, 11, 12, 10, 14, 11, 9, 11, 11, 11, 14, 11, 14, 12, 13, 14, 10, 10, 16, 14, 11, 12, 14, 12, 14, 10, 16], start=1)})

FWD = "{c}{r}*(1+Rate_B*Tenor_Days/Basis_B)/(1+Rate_A*Tenor_Days/Basis_A)"
put(sm, "A6", 0, nf="0")
put(sm, "B6", "-", align=Alignment(horizontal="center"))
put(sm, "C6", "-", align=Alignment(horizontal="center"))
put(sm, "D6", "=Spot_Init", f_link, NF_PX)
put(sm, "E6", "=" + FWD.format(c="D", r=6), f_base, NF_PX)
put(sm, "F6", "=(E6-D6)/Tick", f_base, NF_PIP)
put(sm, "Y6", "=Start_Pos", f_link, NF_QTY)
put(sm, "AA6", 0, f_base, NF_PNL)
put(sm, "AB6", 0, f_base, "0")
put(sm, "AC6", "=Y6*E6", f_base, NF_PNL)
for col in range(1, 30):
    sm.cell(row=6, column=col).fill = fill_grey

zs, actions = generate_scenario(SEED, N_STEPS)
for k in range(N_STEPS):
    r, p = FIRST + k, FIRST + k - 1
    f = {
        "A": (f"=A{p}+1", "0"),
        "D": (f"=D{p}+B{r}*Step_Vol_Ticks*Tick", NF_PX),
        "E": ("=" + FWD.format(c="D", r=r), NF_PX),
        "F": (f"=(E{r}-D{r})/Tick", NF_PIP),
        "G": (f"=Y{p}", NF_QTY),
        "H": (f"=G{r}/Pos_Limit", NF_PCT),
        "I": (f"=-MAX(-1,MIN(1,H{r}))*Max_Skew_Ticks", NF_TICK),
        "J": (f"=ROUND(E{r}/Tick+I{r}-Spread_Ticks/2,0)*Tick", NF_PX),
        "K": (f"=J{r}+Spread_Ticks*Tick", NF_PX),
        "L": (f'=IF(AB{p}=1,"HALTED","LIVE")', None),
        "M": (f'=IF(L{r}="LIVE",IF(C{r}="BUY",-Clip_Size,IF(C{r}="SELL",Clip_Size,0)),0)', NF_QTY),
        "N": (f"=IF(M{r}<0,K{r},IF(M{r}>0,J{r},0))", NF_PX),
        "O": (f"=G{r}+M{r}", NF_QTY),
        "P": (f"=IF(M{r}<0,(K{r}-E{r})*ABS(M{r}),IF(M{r}>0,(E{r}-J{r})*M{r},0))", NF_PNL),
        "Q": (f"=G{r}*(E{r}-E{p})", NF_PNL),
        "R": (f"=AA{p}+Q{r}+P{r}", NF_PNL),
        "S": (f"=IF(ABS(O{r})>Pos_Limit,1,0)", "0;-0;-"),
        "T": (f"=IF(AND(AB{p}=0,R{r}<=-Loss_Limit),1,0)", "0;-0;-"),
        "U": (f'=IF(T{r}=1,"LOSS LIMIT",IF(S{r}=1,"POSITION LIMIT",""))', None),
        "V": (f"=IF(T{r}=1,-O{r},IF(S{r}=1,-(O{r}-SIGN(O{r})*Hedge_Target),0))", NF_QTY),
        "W": (f"=IF(V{r}>0,E{r}+Hedge_Cost_Ticks*Tick,IF(V{r}<0,E{r}-Hedge_Cost_Ticks*Tick,0))", NF_PX),
        "X": (f"=-ABS(V{r})*Hedge_Cost_Ticks*Tick", NF_PNL),
        "Y": (f"=O{r}+V{r}", NF_QTY),
        "Z": (f"=Q{r}+P{r}+X{r}", NF_PNL),
        "AA": (f"=AA{p}+Z{r}", NF_PNL),
        "AB": (f"=IF(Halt_On_Loss=1,MAX(AB{p},T{r}),0)", "0"),
        "AC": (f"=Y{r}*E{r}", NF_PNL),
    }
    for col, (fml, nf) in f.items():
        put(sm, f"{col}{r}", fml, f_base, nf, border=box,
            align=Alignment(horizontal="center") if col in ("L", "U") else None)
    put(sm, f"B{r}", zs[k], f_input, "0.0;-0.0;0.0", fill_input, Alignment(horizontal="center"), box)
    put(sm, f"C{r}", actions[k], f_input, None, fill_input, Alignment(horizontal="center"), box)

dv2 = DataValidation(type="list", formula1='"BUY,SELL,NONE"', allow_blank=False,
                     showErrorMessage=True, errorTitle="Client action", error="Choose BUY, SELL or NONE.")
sm.add_data_validation(dv2)
dv2.add(f"C{FIRST}:C{LAST}")
dv3 = DataValidation(type="decimal", operator="between", formula1="-8", formula2="8",
                     showErrorMessage=True, errorTitle="Market move", error="Enter a move in sigma units between -8 and 8.")
sm.add_data_validation(dv3)
dv3.add(f"B{FIRST}:B{LAST}")

rng = f"A{FIRST}:AC{LAST}"
sm.conditional_formatting.add(rng, FormulaRule(formula=[f"$T{FIRST}=1"], fill=PatternFill("solid", bgColor="F4B084"), stopIfTrue=True))
sm.conditional_formatting.add(rng, FormulaRule(formula=[f"$S{FIRST}=1"], fill=PatternFill("solid", bgColor="FFE699"), stopIfTrue=True))
sm.conditional_formatting.add(rng, FormulaRule(formula=[f'$L{FIRST}="HALTED"'], font=Font(color="808080"),
                                               fill=PatternFill("solid", bgColor="EDEDED")))
sm.freeze_panes = "D6"

# --------------------------------------------------------------------------------------- Summary
su = wb.create_sheet("Summary")
su.sheet_view.showGridLines = False
widths(su, {"A": 2, "B": 46, "C": 18, "D": 56})
put(su, "B1", '="Simulation Summary: "&Instrument', f_title)
put(su, "B2", "All figures are live formulas over the Simulation sheet.", f_sub)
S = lambda col: f"Simulation!${col}${FIRST}:${col}${LAST}"
sections = [
    ("ACTIVITY", [
        ("Steps simulated", f"=COUNT({S('A')})", "0", ""),
        ("Client fills", f'=COUNTIF({S("M")},"<>0")', "0", "RFQs that traded (quoting live)"),
        ("  Client bought (we sold at Ask)", f'=COUNTIF({S("M")},"<0")', "0", ""),
        ("  Client sold (we bought at Bid)", f'=COUNTIF({S("M")},">0")', "0", ""),
        ('="Client volume traded ("&Unit_Label&")"', f"=SUMPRODUCT(ABS({S('M')}))", NF_QTY, ""),
        ("Position-limit hedges", f"=SUMPRODUCT(({S('S')}=1)*({S('T')}=0))", "0", "Hedged back to the residual level"),
        ("Loss-limit flattenings", f"=SUM({S('T')})", "0", ""),
        ('="Hedge volume traded with street ("&Unit_Label&")"', f"=SUMPRODUCT(ABS({S('V')}))", NF_QTY, ""),
    ]),
    ("RISK USAGE", [
        ('="Largest |position| after client fill ("&Unit_Label&")"', f"=MAX(MAX({S('O')}),-MIN({S('O')}))", NF_QTY, "Before any hedge"),
        ('="Largest |closing position| ("&Unit_Label&")"', f"=MAX(MAX({S('Y')}),-MIN({S('Y')}))", NF_QTY, "After hedging"),
        ('="Position limit ("&Unit_Label&")"', "=Pos_Limit", NF_QTY, ""),
        ('="Worst cumulative P&L ("&PnL_Ccy&")"', f"=MIN(Simulation!$AA$6:$AA${LAST})", NF_PNL, "@worst"),
        ("Loss limit used at worst point", "=MAX(0,-C{worst})/Loss_Limit", NF_PCT, "100% = limit breached"),
        ("Quoting halted at end of run?", f'=IF(Simulation!$AB${LAST}=1,"YES","NO")', None, ""),
        ("Step at which loss limit hit", f'=IFERROR(INDEX({S("A")},MATCH(1,{S("T")},0)),"n/a")', "0", ""),
    ]),
    ('="P&L ATTRIBUTION ("&PnL_Ccy&")"', [
        ("Spread capture vs fair forward", f"=SUM({S('P')})", NF_PNL, "@a"),
        ("Mark-to-market on inventory", f"=SUM({S('Q')})", NF_PNL, "Opening position x change in fair forward"),
        ("Hedge execution cost", f"=SUM({S('X')})", NF_PNL, "@b"),
        ("Total P&L", "=SUM(C{a}:C{b})", NF_PNL, "@tot"),
        ("Closing cumulative P&L", f"=Simulation!$AA${LAST}", NF_PNL, "@cum"),
    ]),
    ("INTEGRITY CHECKS", [
        ("P&L attribution reconciles", '=IF(ABS(C{tot}-C{cum})<0.005,"OK","ERROR")', None, "Components sum to cumulative P&L"),
        ("Position roll-forward reconciles",
         f'=IF(ABS(Start_Pos+SUM({S("M")})+SUM({S("V")})-Simulation!$Y${LAST})<0.5,"OK","ERROR")', None,
         "Start + client fills + hedges = closing position"),
        ("Hedge respects limit",
         f'=IF(SUMPRODUCT(({S("Y")}<>0)*(ABS({S("Y")})>Pos_Limit))=0,"OK","ERROR")', None,
         "No step ends outside the position limit"),
    ]),
]
r = 4
anchors = {}
for title, items in sections:
    band(su, r, title, "BCD")
    r += 1
    for lab, fml, nf, note in items:
        if lab.startswith("Worst"):
            anchors["worst"] = r
        elif lab.startswith('="Worst'):
            anchors["worst"] = r
        if lab == "Spread capture vs fair forward":
            anchors["a"] = r
        if lab == "Hedge execution cost":
            anchors["b"] = r
        if lab == "Total P&L":
            anchors["tot"] = r
        if lab == "Closing cumulative P&L":
            anchors["cum"] = r
        if "{" in fml:
            fml = fml.format(**anchors)
        note = "" if note.startswith("@") else note
        is_tot = lab == "Total P&L"
        put(su, f"B{r}", lab, f_bold if is_tot else f_base)
        put(su, f"C{r}", fml, f_bold if is_tot else f_base, nf, align=Alignment(horizontal="right"))
        put(su, f"D{r}", note, f_sub, align=Alignment(indent=1))
        r += 1
    r += 1

ok_range = f"C{r - 4}:C{r - 2}"
su.conditional_formatting.add(ok_range, FormulaRule(formula=[f'C{r - 4}="OK"'], font=Font(name=FONT, color="008000", bold=True)))
su.conditional_formatting.add(ok_range, FormulaRule(formula=[f'C{r - 4}="ERROR"'], font=Font(name=FONT, color="FF0000", bold=True)))

cats = Reference(sm, min_col=1, min_row=6, max_row=LAST)
for title, col, anchor, ytitle in [("Cumulative P&L", 27, "F4", "P&L"), ("Closing position (units)", 25, "F22", "Units")]:
    ch = LineChart()
    ch.title = title
    ch.height, ch.width = 8.5, 17
    ch.y_axis.title = ytitle
    ch.x_axis.title = "Step"
    ch.add_data(Reference(sm, min_col=col, min_row=5, max_row=LAST), titles_from_data=True)
    ch.set_categories(cats)
    ch.legend = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.tickLblSkip = 10
    ch.x_axis.tickLblPos = "low"
    ch.series[0].smooth = False
    ch.series[0].graphicalProperties.line.width = 20000
    su.add_chart(ch, anchor)

# sheet order: Guide, Inputs, Instruments, Live, Pricing, Simulation, Summary
wb.move_sheet("Inputs", offset=1 - wb.index(wb["Inputs"]))
for sheet in wb.worksheets:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
sm.print_title_rows = "4:5"
wb.save(_ARGS.out)
print("saved", _ARGS.out, "|", [s.title for s in wb.worksheets])
