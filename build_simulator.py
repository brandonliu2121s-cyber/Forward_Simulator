"""Builds Forwards_Market_Making_Simulator.xlsx (USD/SGD forwards market-making simulation).

Every number the workbook computes is a live Excel formula driven by the Inputs sheet and the
two blue per-step input columns on the Simulation sheet (spot shock, client action).
Run:  python3 build_simulator.py   then   python3 <skill>/scripts/recalc.py Forwards_Market_Making_Simulator.xlsx
"""
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

from reference_model import DEFAULTS as D, generate_scenario

OUT = "Forwards_Market_Making_Simulator.xlsx"
SEED, N_STEPS = 277, 60

FONT = "Arial"
NAVY, GREY, LIGHT = "1F3864", "F2F2F2", "D9E1F2"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_sub = Font(name=FONT, size=10, italic=True, color="595959")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_group = PatternFill("solid", fgColor=LIGHT)
fill_input = PatternFill("solid", fgColor="FFFF99")
fill_grey = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

NF_PX = "0.0000"
NF_USD = "#,##0;(#,##0);-"
NF_SGD = "#,##0.00;(#,##0.00);-"
NF_PCT = "0.0%;(0.0%);-"
NF_TICK = "0.00;(0.00);-"
NF_PIP = "0.0;(0.0);-"

wb = Workbook()


def put(ws, ref, value, font=f_base, nf=None, fill=None, align=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if nf:
        c.number_format = nf
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


# --------------------------------------------------------------------------------------- Guide
g = wb.active
g.title = "Guide"
g.sheet_view.showGridLines = False
widths(g, {"A": 2, "B": 26, "C": 100})
put(g, "B1", "USD/SGD Forwards Market-Making Simulator", f_title)
put(g, "B2", "Simulated dealing desk: price a 45-day USD/SGD forward off interest-rate parity, quote USD 250k a side, "
             "skew quotes with inventory, and hedge back-to-back when position or loss limits are breached.", f_sub)
rows = [
    ("SHEETS", None),
    ("Inputs", "Every assumption in one place (blue on yellow = editable). Change anything and the whole model re-prices."),
    ("Pricing", "Step-by-step covered-interest-parity derivation of the fair forward, a live quote calculator for any "
                "inventory, and a skew ladder showing quotes at every position."),
    ("Simulation", "60-step market replay. Edit the two blue columns (spot shock in ticks, client action BUY/SELL/NONE) "
                   "to run your own scenario; everything else is formulas."),
    ("Summary", "Fills, hedges, P&L attribution (spread capture / mark-to-market / hedge cost), limit usage, "
                "reconciliation checks and charts."),
    ("PRICING", None),
    ("Fair forward", "F = Spot x (1 + r_quote x days / basis_quote) / (1 + r_base x days / basis_base). "
                     "USD is the base currency (4.00%), SGD the quoted currency (0.55%), tenor 45 days. "
                     "USD yields more than SGD, so USD trades at a forward discount (negative forward points)."),
    ("Quotes", "Mid = fair forward + skew. Bid = mid - spread/2, Ask = Bid + spread (5 ticks, 1 tick = 0.0001), "
               "rounded to the tick. Each quote is for USD 250,000."),
    ("Skew", "Skew (ticks) = -(position / position limit, capped at +/-100%) x max skew. Long USD pushes both "
             "quotes down (cheap to buy from us, unattractive to sell to us); short USD pushes them up."),
    ("RISK", None),
    ("Position limit", "If |position| after a client fill exceeds the limit, a back-to-back hedge trades the street "
                       "to bring |position| down to the hedge-to level, paying the street half-spread."),
    ("Loss limit", "If cumulative P&L (before the hedge cost of that step) falls to -loss limit, the book is "
                   "flattened with the street. With 'halt' = 1 quoting then stops for the rest of the run."),
    ("CONVENTIONS", None),
    ("Position", "Net USD bought forward from clients. + = long USD / short SGD. Client BUY = client buys USD from us "
                 "at our Ask (we go shorter); client SELL = client sells USD to us at our Bid."),
    ("P&L", "In SGD, measured against fair forward: spread capture (fill vs fair) + mark-to-market of opening "
            "position on the fair-forward move + street hedge cost."),
    ("LIMITATIONS", None),
    ("Simplifications", "Constant 45-day tenor each step (no roll-down/carry decay); P&L not discounted; one RFQ per "
                        "step at a fixed size; hedges fill instantly at fair +/- street half-spread; spot is an "
                        "illustrative level, not live market data; the default scenario is randomly generated "
                        "(seed %d) and for demonstration only." % SEED),
]
r = 4
for k, v in rows:
    if v is None:
        put(g, f"B{r}", k, f_head, fill=fill_head)
        put(g, f"C{r}", None, f_head, fill=fill_head)
    else:
        put(g, f"B{r}", k, f_bold, align=Alignment(vertical="top"))
        put(g, f"C{r}", v, align=Alignment(wrap_text=True, vertical="top"))
        g.row_dimensions[r].height = 15 * (1 + len(v) // 105)
    r += 1

# --------------------------------------------------------------------------------------- Inputs
ws = wb.create_sheet("Inputs")
ws.sheet_view.showGridLines = False
widths(ws, {"A": 2, "B": 34, "C": 16, "D": 12, "E": 90})
put(ws, "B1", "Inputs & Assumptions", f_title)
put(ws, "B2", "Blue text on yellow = editable input. All other cells in the workbook are formulas.", f_sub)
for col, h in zip("BCDE", ["Assumption", "Value", "Unit", "Source / note"]):
    put(ws, f"{col}4", h, f_head, fill=fill_head)

USER = "Specified in project brief"
inputs = [
    ("MARKET & RATES", None),
    ("Spot_Init", "Starting USD/SGD spot", D["spot_init"], "SGD per USD", NF_PX,
     "Illustrative level (approx. Aug 2025); not live market data. Overwrite with a current spot."),
    ("Base_IR", "Base currency (USD) interest rate", D["base_ir"], "% p.a.", "0.00%", USER + " (Base IR 4.00%)."),
    ("Quote_IR", "Quoted currency (SGD) interest rate", D["quote_ir"], "% p.a.", "0.00%", USER + " (Quoted IR 0.55%)."),
    ("Tenor_Days", "Forward tenor", D["tenor_days"], "days", "0", USER + " (45-day tenor)."),
    ("Base_Basis", "USD day-count basis", D["base_basis"], "days", "0", "Market convention ACT/360 for USD money-market rates."),
    ("Quote_Basis", "SGD day-count basis", D["quote_basis"], "days", "0", "Market convention ACT/365 for SGD rates."),
    ("QUOTING", None),
    ("Tick", "Tick size", D["tick"], "SGD per USD", "0.0000", "USD/SGD is quoted to 4 decimals; 1 tick = 1 pip."),
    ("Clip_Size", "Quote size per side", D["clip"], "USD", NF_USD, USER + " (USD 250k per side)."),
    ("Spread_Ticks", "Bid/offer spread", D["spread_ticks"], "ticks", "0.0", USER + " (5-tick spread)."),
    ("Max_Skew_Ticks", "Max skew at full position limit", D["max_skew_ticks"], "ticks", "0.0",
     "Assumption: how hard quotes lean away from inventory when position = limit. Scales linearly below that."),
    ("RISK LIMITS & HEDGING", None),
    ("Pos_Limit", "Position limit (absolute)", D["pos_limit"], "USD", NF_USD,
     "Assumption (4 clips). Breach -> back-to-back hedge with the street."),
    ("Hedge_Target", "Residual position after position-limit hedge", D["hedge_target"], "USD", NF_USD,
     "Assumption: 0 = hedge fully flat; e.g. 250,000 keeps a one-clip residual in the same direction."),
    ("Hedge_Cost_Ticks", "Street half-spread paid on hedges", D["hedge_cost_ticks"], "ticks", "0.0",
     "Assumption: hedges execute at fair forward -/+ this many ticks."),
    ("Loss_Limit", "Cumulative loss limit", D["loss_limit"], "SGD", NF_USD,
     "Assumption. Breach -> flatten the book with the street."),
    ("Halt_On_Loss", "Halt quoting after loss-limit breach (1 = yes, 0 = no)", D["halt_on_loss"], "flag", "0",
     "Assumption. 0 = keep quoting (the desk re-flattens every step cumulative P&L stays below the limit)."),
    ("SIMULATION", None),
    ("Start_Pos", "Opening USD position", D["start_pos"], "USD", NF_USD, "Assumption: start the day flat."),
]
r = 5
for item in inputs:
    if item[1] is None:
        put(ws, f"B{r}", item[0], f_bold, fill=fill_group)
        for col in "CDE":
            put(ws, f"{col}{r}", None, f_bold, fill=fill_group)
    else:
        nm, label, val, unit, nf, note = item
        put(ws, f"B{r}", label)
        put(ws, f"C{r}", val, f_input, nf, fill_input, Alignment(horizontal="right"), box)
        put(ws, f"D{r}", unit, f_sub)
        put(ws, f"E{r}", note, f_sub)
        name(nm, f"Inputs!$C${r}")
    r += 1

dv = DataValidation(type="list", formula1='"0,1"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(ws[f"C{[i for i, it in enumerate(inputs) if it[0] == 'Halt_On_Loss'][0] + 5}"])

# --------------------------------------------------------------------------------------- Pricing
pr = wb.create_sheet("Pricing")
pr.sheet_view.showGridLines = False
widths(pr, {"A": 2, "B": 36, "C": 16, "D": 14, "E": 14, "F": 14, "G": 22, "H": 4})
put(pr, "B1", "Pricing: Interest-Rate Parity & Quote Calculator", f_title)
put(pr, "B2", "Fair forward from covered interest parity, then the skewed two-way quote for any inventory.", f_sub)

put(pr, "B4", "1. Fair forward (covered interest parity)", f_head, fill=fill_head)
for col in "CD":
    put(pr, f"{col}4", None, f_head, fill=fill_head)
calc = [
    ("Spot (SGD per USD)", "=Spot_Init", NF_PX, f_link),
    ("USD growth factor  1 + r_USD x d / 360", "=1+Base_IR*Tenor_Days/Base_Basis", "0.000000", f_base),
    ("SGD growth factor  1 + r_SGD x d / 365", "=1+Quote_IR*Tenor_Days/Quote_Basis", "0.000000", f_base),
    ("Fair forward outright", "=C5*C7/C6", "0.00000", f_bold),
    ("Forward points (ticks)", "=(C8-C5)/Tick", NF_PIP, f_base),
    ("Forward premium / (discount) p.a.", "=(C8/C5-1)*365/Tenor_Days", "0.00%;(0.00%)", f_base),
    ("Implied rate differential (SGD - USD)", "=Quote_IR-Base_IR", "0.00%;(0.00%)", f_base),
]
for i, (lab, fml, nf, fnt) in enumerate(calc):
    put(pr, f"B{5 + i}", lab)
    put(pr, f"C{5 + i}", fml, fnt, nf)
put(pr, "D9", "USD at forward discount: USD yield > SGD yield", f_sub)

put(pr, "B13", "2. Live quote calculator", f_head, fill=fill_head)
for col in "CD":
    put(pr, f"{col}13", None, f_head, fill=fill_head)
put(pr, "B14", "Current USD position (edit me)")
put(pr, "C14", 500000, f_input, NF_USD, fill_input, border=box)
put(pr, "D14", "+ = long USD", f_sub)
q = [
    ("Position utilisation of limit", "=C14/Pos_Limit", NF_PCT),
    ("Skew (ticks)", "=-MAX(-1,MIN(1,C15))*Max_Skew_Ticks", NF_TICK),
    ("Bid  (client sells USD to us)", "=ROUND(C8/Tick+C16-Spread_Ticks/2,0)*Tick", NF_PX),
    ("Ask  (client buys USD from us)", "=C17+Spread_Ticks*Tick", NF_PX),
    ("Quoted spread (ticks)", "=(C18-C17)/Tick", NF_TICK),
    ("Mid vs fair forward (ticks)", "=((C17+C18)/2-C8)/Tick", NF_TICK),
    ("Quote size each side (USD)", "=Clip_Size", NF_USD),
    ("Position after being lifted on the Ask", "=C14-Clip_Size", NF_USD),
    ("Position after being hit on the Bid", "=C14+Clip_Size", NF_USD),
    ("Next fill triggers hedge?", '=IF(OR(ABS(C21)>Pos_Limit,ABS(C22)>Pos_Limit),"YES: a fill can breach the limit","No")', None),
]
for i, (lab, fml, nf) in enumerate(q):
    put(pr, f"B{15 + i}", lab)
    put(pr, f"C{15 + i}", fml, f_bold if i in (2, 3) else f_base, nf, align=Alignment(horizontal="right"))

put(pr, "B27", "3. Skew ladder: quotes at every inventory level", f_head, fill=fill_head)
for col in "CDEFG":
    put(pr, f"{col}27", None, f_head, fill=fill_head)
for col, h in zip("BCDEFG", ["Position (USD)", "Utilisation", "Skew (ticks)", "Bid", "Ask", "Limit check"]):
    put(pr, f"{col}28", h, f_bold, fill=fill_group, align=Alignment(horizontal="center", wrap_text=True))
for i in range(13):
    rr = 29 + i
    put(pr, f"B{rr}", "=-6*Clip_Size" if i == 0 else f"=B{rr - 1}+Clip_Size", f_base, NF_USD)
    put(pr, f"C{rr}", f"=B{rr}/Pos_Limit", f_base, NF_PCT)
    put(pr, f"D{rr}", f"=-MAX(-1,MIN(1,C{rr}))*Max_Skew_Ticks", f_base, NF_TICK)
    put(pr, f"E{rr}", f"=ROUND($C$8/Tick+D{rr}-Spread_Ticks/2,0)*Tick", f_base, NF_PX)
    put(pr, f"F{rr}", f"=E{rr}+Spread_Ticks*Tick", f_base, NF_PX)
    put(pr, f"G{rr}", f'=IF(ABS(B{rr})>Pos_Limit,"BREACH: hedge","within limit")', f_base, align=Alignment(indent=1))
pr.conditional_formatting.add("B29:G41", FormulaRule(formula=["ABS($B29)>Pos_Limit"],
                              fill=PatternFill("solid", bgColor="F8CBAD")))

# --------------------------------------------------------------------------------------- Simulation
sm = wb.create_sheet("Simulation")
sm.sheet_view.showGridLines = False
FIRST, LAST = 7, 6 + N_STEPS  # step 0 lives on row 6
put(sm, "A1", "Market-Making Simulation", f_title)
put(sm, "A2", "One client RFQ per step. Edit the blue columns B:C to run your own scenario; all other cells are formulas. "
              "Position = net USD bought from clients (+ long).", f_sub)
put(sm, "A3", "Default scenario: spot shocks and client flow randomly generated by build_simulator.py (seed %d); "
              "illustrative only." % SEED, f_sub)

groups = [("A", "C", "Scenario input"), ("D", "F", "Market"), ("G", "I", "Exposure & skew"), ("J", "L", "Our quote"),
          ("M", "P", "Client trade"), ("Q", "T", "P&L and limit checks"), ("U", "X", "Back-to-back hedge"),
          ("Y", "AC", "Close of step")]
for a, b, label in groups:
    sm.merge_cells(f"{a}4:{b}4")
    put(sm, f"{a}4", label, f_bold, fill=fill_group, align=Alignment(horizontal="center"))
heads = [
    "Step", "Spot shock (ticks)", "Client action\n(BUY / SELL / NONE)",
    "Spot", "Fair forward", "Fwd points (ticks)",
    "Opening position (USD)", "Utilisation of limit", "Skew (ticks)",
    "Bid", "Ask", "Quoting status",
    "Client fill (USD, + = we buy)", "Fill price", "Position after client trade", "Spread capture (SGD)",
    "MTM on opening position (SGD)", "Cum. P&L before hedge (SGD)", "Position limit breach", "Loss limit breach",
    "Hedge trigger", "Hedge qty (USD, + = we buy)", "Hedge price", "Hedge cost (SGD)",
    "Closing position (USD)", "Step P&L (SGD)", "Cumulative P&L (SGD)", "Quoting halted (1 = yes)",
    "Closing position in SGD",
]
for i, h in enumerate(heads, start=1):
    put(sm, f"{get_column_letter(i)}5", h, f_head, fill=fill_head,
        align=Alignment(horizontal="center", vertical="center", wrap_text=True), border=box)
sm.row_dimensions[5].height = 54
widths(sm, {get_column_letter(i): w for i, w in enumerate(
    [6, 10, 15, 9, 11, 10, 14, 11, 9, 9, 9, 11, 14, 9, 14, 12, 13, 14, 10, 10, 16, 14, 9, 12, 14, 12, 14, 10, 16], start=1)})

FWD = "{c}{r}*(1+Quote_IR*Tenor_Days/Quote_Basis)/(1+Base_IR*Tenor_Days/Base_Basis)"
# step 0 (opening state)
put(sm, "A6", 0, nf="0")
put(sm, "B6", "-", align=Alignment(horizontal="center"))
put(sm, "C6", "-", align=Alignment(horizontal="center"))
put(sm, "D6", "=Spot_Init", f_link, NF_PX)
put(sm, "E6", "=" + FWD.format(c="D", r=6), f_base, NF_PX)
put(sm, "F6", "=(E6-D6)/Tick", f_base, NF_PIP)
put(sm, "Y6", "=Start_Pos", f_link, NF_USD)
put(sm, "AA6", 0, f_base, NF_SGD)
put(sm, "AB6", 0, f_base, "0")
put(sm, "AC6", "=Y6*E6", f_base, NF_SGD)
for col in range(1, 30):
    sm.cell(row=6, column=col).fill = fill_grey

shocks, actions = generate_scenario(SEED, N_STEPS)
for k in range(N_STEPS):
    r, p = FIRST + k, FIRST + k - 1
    f = {
        "A": (f"=A{p}+1", "0"),
        "D": (f"=D{p}+B{r}*Tick", NF_PX),
        "E": ("=" + FWD.format(c="D", r=r), NF_PX),
        "F": (f"=(E{r}-D{r})/Tick", NF_PIP),
        "G": (f"=Y{p}", NF_USD),
        "H": (f"=G{r}/Pos_Limit", NF_PCT),
        "I": (f"=-MAX(-1,MIN(1,H{r}))*Max_Skew_Ticks", NF_TICK),
        "J": (f"=ROUND(E{r}/Tick+I{r}-Spread_Ticks/2,0)*Tick", NF_PX),
        "K": (f"=J{r}+Spread_Ticks*Tick", NF_PX),
        "L": (f'=IF(AB{p}=1,"HALTED","LIVE")', None),
        "M": (f'=IF(L{r}="LIVE",IF(C{r}="BUY",-Clip_Size,IF(C{r}="SELL",Clip_Size,0)),0)', NF_USD),
        "N": (f"=IF(M{r}<0,K{r},IF(M{r}>0,J{r},0))", NF_PX),
        "O": (f"=G{r}+M{r}", NF_USD),
        "P": (f"=IF(M{r}<0,(K{r}-E{r})*ABS(M{r}),IF(M{r}>0,(E{r}-J{r})*M{r},0))", NF_SGD),
        "Q": (f"=G{r}*(E{r}-E{p})", NF_SGD),
        "R": (f"=AA{p}+Q{r}+P{r}", NF_SGD),
        "S": (f"=IF(ABS(O{r})>Pos_Limit,1,0)", "0;-0;-"),
        "T": (f"=IF(AND(AB{p}=0,R{r}<=-Loss_Limit),1,0)", "0;-0;-"),
        "U": (f'=IF(T{r}=1,"LOSS LIMIT",IF(S{r}=1,"POSITION LIMIT",""))', None),
        "V": (f"=IF(T{r}=1,-O{r},IF(S{r}=1,-(O{r}-SIGN(O{r})*Hedge_Target),0))", NF_USD),
        "W": (f"=IF(V{r}>0,E{r}+Hedge_Cost_Ticks*Tick,IF(V{r}<0,E{r}-Hedge_Cost_Ticks*Tick,0))", NF_PX),
        "X": (f"=-ABS(V{r})*Hedge_Cost_Ticks*Tick", NF_SGD),
        "Y": (f"=O{r}+V{r}", NF_USD),
        "Z": (f"=Q{r}+P{r}+X{r}", NF_SGD),
        "AA": (f"=AA{p}+Z{r}", NF_SGD),
        "AB": (f"=IF(Halt_On_Loss=1,MAX(AB{p},T{r}),0)", "0"),
        "AC": (f"=Y{r}*E{r}", NF_SGD),
    }
    for col, (fml, nf) in f.items():
        put(sm, f"{col}{r}", fml, f_base, nf, border=box,
            align=Alignment(horizontal="center") if col in ("L", "U") else None)
    put(sm, f"B{r}", shocks[k], f_input, "0;-0;0", fill_input, Alignment(horizontal="center"), box)
    put(sm, f"C{r}", actions[k], f_input, None, fill_input, Alignment(horizontal="center"), box)

dv2 = DataValidation(type="list", formula1='"BUY,SELL,NONE"', allow_blank=False,
                     showErrorMessage=True, errorTitle="Client action", error="Choose BUY, SELL or NONE.")
sm.add_data_validation(dv2)
dv2.add(f"C{FIRST}:C{LAST}")
dv3 = DataValidation(type="whole", operator="between", formula1="-1000", formula2="1000",
                     showErrorMessage=True, errorTitle="Spot shock", error="Whole number of ticks.")
sm.add_data_validation(dv3)
dv3.add(f"B{FIRST}:B{LAST}")

rng = f"A{FIRST}:AC{LAST}"
sm.conditional_formatting.add(rng, FormulaRule(formula=[f"$T{FIRST}=1"], fill=PatternFill("solid", bgColor="F4B084"), stopIfTrue=True))
sm.conditional_formatting.add(rng, FormulaRule(formula=[f"$S{FIRST}=1"], fill=PatternFill("solid", bgColor="FFE699"), stopIfTrue=True))
sm.conditional_formatting.add(rng, FormulaRule(formula=[f'$L{FIRST}="HALTED"'], font=Font(color="808080"),
                                               fill=PatternFill("solid", bgColor="EDEDED")))
sm.freeze_panes = "D6"

# --------------------------------------------------------------------------------------- Summary
su = wb.create_sheet("Summary")
su.sheet_view.showGridLines = False
widths(su, {"A": 2, "B": 44, "C": 18, "D": 56})
put(su, "B1", "Simulation Summary", f_title)
put(su, "B2", "All figures are live formulas over the Simulation sheet.", f_sub)
S = lambda col: f"Simulation!${col}${FIRST}:${col}${LAST}"
sections = [
    ("ACTIVITY", [
        ("Steps simulated", f"=COUNT({S('A')})", "0", ""),
        ("Client fills", f'=COUNTIF({S("M")},"<>0")', "0", "RFQs that traded (quoting live)"),
        ("  Client bought USD (we sold at Ask)", f'=COUNTIF({S("M")},"<0")', "0", ""),
        ("  Client sold USD (we bought at Bid)", f'=COUNTIF({S("M")},">0")', "0", ""),
        ("Client notional traded (USD)", f"=SUMPRODUCT(ABS({S('M')}))", NF_USD, ""),
        ("Position-limit hedges", f"=SUMPRODUCT(({S('S')}=1)*({S('T')}=0))", "0", "Hedged back to the hedge-to level"),
        ("Loss-limit flattenings", f"=SUM({S('T')})", "0", ""),
        ("Hedge notional traded with street (USD)", f"=SUMPRODUCT(ABS({S('V')}))", NF_USD, ""),
    ]),
    ("RISK USAGE", [
        ("Largest |position| after client fill (USD)", f"=MAX(MAX({S('O')}),-MIN({S('O')}))", NF_USD, "Before any hedge"),
        ("Largest |closing position| (USD)", f"=MAX(MAX({S('Y')}),-MIN({S('Y')}))", NF_USD, "After hedging"),
        ("Position limit (USD)", "=Pos_Limit", NF_USD, ""),
        ("Worst cumulative P&L (SGD)", f"=MIN(Simulation!$AA$6:$AA${LAST})", NF_SGD, ""),
        ("Loss limit used at worst point", f"=MAX(0,-C{{r_worst}})/Loss_Limit", NF_PCT, "100% = limit breached"),
        ("Quoting halted at end of run?", f'=IF(Simulation!$AB${LAST}=1,"YES","NO")', None, ""),
        ("Step at which loss limit hit", f'=IFERROR(INDEX({S("A")},MATCH(1,{S("T")},0)),"n/a")', "0", ""),
    ]),
    ("P&L ATTRIBUTION (SGD)", [
        ("Spread capture vs fair forward", f"=SUM({S('P')})", NF_SGD, "Edge earned on client fills"),
        ("Mark-to-market on inventory", f"=SUM({S('Q')})", NF_SGD, "Opening position x change in fair forward"),
        ("Hedge execution cost", f"=SUM({S('X')})", NF_SGD, "Street half-spread paid on hedges"),
        ("Total P&L", "=SUM(C{a}:C{b})", NF_SGD, ""),
        ("Closing cumulative P&L", f"=Simulation!$AA${LAST}", NF_SGD, ""),
    ]),
    ("INTEGRITY CHECKS", [
        ("P&L attribution reconciles", '=IF(ABS(C{tot}-C{cum})<0.005,"OK","ERROR")', None, "Components sum to cumulative P&L"),
        ("Position roll-forward reconciles",
         f'=IF(ABS(Start_Pos+SUM({S("M")})+SUM({S("V")})-Simulation!$Y${LAST})<0.5,"OK","ERROR")', None,
         "Start + client fills + hedges = closing position"),
        ("Hedge respects limit",
         f'=IF(SUMPRODUCT(({S("Y")}<>0)*(ABS({S("Y")})>Pos_Limit))=0,"OK","ERROR")', None,
         "No step ends outside the position limit"),
    ]),
]
r = 4
anchors = {}
for title, items in sections:
    put(su, f"B{r}", title, f_head, fill=fill_head)
    put(su, f"C{r}", None, f_head, fill=fill_head)
    put(su, f"D{r}", None, f_head, fill=fill_head)
    r += 1
    start = r
    for lab, fml, nf, note in items:
        if "{r_worst}" in fml:
            fml = fml.replace("{r_worst}", str(anchors["worst"]))
        if lab == "Worst cumulative P&L (SGD)":
            anchors["worst"] = r
        if lab == "Spread capture vs fair forward":
            anchors["a"] = r
        if lab == "Hedge execution cost":
            anchors["b"] = r
        if lab == "Total P&L":
            fml = fml.format(a=anchors["a"], b=anchors["b"])
            anchors["tot"] = r
        if lab == "Closing cumulative P&L":
            anchors["cum"] = r
        if "{tot}" in fml:
            fml = fml.format(tot=anchors["tot"], cum=anchors["cum"])
        put(su, f"B{r}", lab, f_bold if lab == "Total P&L" else f_base)
        put(su, f"C{r}", fml, f_bold if lab == "Total P&L" else f_base, nf, align=Alignment(horizontal="right"))
        put(su, f"D{r}", note, f_sub, align=Alignment(indent=1))
        r += 1
    r += 1

ok_range = f"C{r - 4}:C{r - 2}"
su.conditional_formatting.add(ok_range, FormulaRule(formula=[f'C{r - 4}="OK"'], font=Font(name=FONT, color="008000", bold=True)))
su.conditional_formatting.add(ok_range, FormulaRule(formula=[f'C{r - 4}="ERROR"'], font=Font(name=FONT, color="FF0000", bold=True)))

cats = Reference(sm, min_col=1, min_row=6, max_row=LAST)
for title, col, anchor, ytitle in [("Cumulative P&L (SGD)", 27, "F4", "SGD"), ("Closing USD position", 25, "F22", "USD")]:
    ch = LineChart()
    ch.title = title
    ch.height, ch.width = 8.5, 17
    ch.y_axis.title = ytitle
    ch.x_axis.title = "Step"
    ch.add_data(Reference(sm, min_col=col, min_row=5, max_row=LAST), titles_from_data=True)
    ch.set_categories(cats)
    ch.legend = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.tickLblSkip = 10
    ch.x_axis.tickLblPos = "low"
    ch.series[0].smooth = False
    ch.series[0].graphicalProperties.line.width = 20000
    su.add_chart(ch, anchor)

for sheet in wb.worksheets:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
sm.print_title_rows = "4:5"
wb.save(OUT)
print("saved", OUT)
